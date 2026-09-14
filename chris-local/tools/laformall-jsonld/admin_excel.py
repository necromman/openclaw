"""고도몰 관리자 상품 목록 엑셀다운로드 파일 읽기.

관리자 > 상품 관리 > 상품 목록 > 엑셀다운로드 로 받은 xlsx·csv 를 넣으면
상품코드로 URL 을 만들고 상품명·판매가·노출/판매 상태를 표에 채운다.
사이트에서 긁어 모으는 것보다 이 파일이 가장 정확하다.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

import fetcher
import fields

# 열 이름은 부분 일치로 찾는다(순서 무관). 앞에 있는 후보를 먼저 본다.
HINTS: dict[str, tuple[str, ...]] = {
    "code": ("상품코드", "goodsno", "goods_no", "상품번호", "코드"),
    "name": ("상품명", "goodsnm", "상품 이름", "제품명"),
    "price": ("판매가", "goodsprice", "판매 가격", "가격"),
    "expose": ("노출상태", "노출여부", "노출", "전시상태", "진열"),
    "sale": ("판매상태", "판매여부", "품절", "재고상태", "판매"),
    "supplier": ("공급사", "공급업체", "브랜드"),
    "date": ("등록일", "수정일", "일자"),
}
REQUIRED = ("code",)
LABELS = {
    "code": "상품코드",
    "name": "상품명",
    "price": "판매가",
    "expose": "노출상태",
    "sale": "판매상태",
    "supplier": "공급사",
    "date": "등록일/수정일",
}

HIDDEN_WORDS = ("미노출", "숨김", "비노출", "미진열", "n", "no", "off", "false", "0")
SHOWN_WORDS = ("노출", "진열", "y", "yes", "on", "true", "1")
SOLDOUT_WORDS = ("품절", "일시품절", "재고없음", "판매중지", "판매종료", "중지")


def read_table(path: str | Path) -> tuple[list[str], list[list[str]]]:
    """xlsx 또는 csv 에서 (헤더, 행 목록) 을 낸다. 헤더 줄을 스스로 찾는다."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        rows = _read_xlsx(path)
    else:
        rows = _read_csv(path)
    rows = [row for row in rows if any(str(cell).strip() for cell in row)]
    if not rows:
        raise RuntimeError("빈 파일입니다")
    head_index = 0
    for index, row in enumerate(rows[:10]):
        joined = " ".join(str(cell) for cell in row)
        if any(hint in joined for hint in ("상품코드", "상품명", "판매가")):
            head_index = index
            break
    headers = [str(cell).strip() for cell in rows[head_index]]
    body = [[str(cell).strip() if cell is not None else "" for cell in row] for row in rows[head_index + 1 :]]
    return headers, body


def _read_xlsx(path: Path) -> list[list]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError(
            f"xlsx 를 읽을 수 없습니다({exc}). 엑셀에서 csv 로 저장한 뒤 다시 시도하세요."
        )
    book = load_workbook(filename=str(path), read_only=True, data_only=True)
    sheet = book.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    book.close()
    return [["" if cell is None else cell for cell in row] for row in rows]


def _read_csv(path: Path) -> list[list]:
    for encoding in ("utf-8-sig", "cp949", "utf-8"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    sample = text[:4000]
    delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    return [row for row in csv.reader(text.splitlines(), delimiter=delimiter)]


def match_columns(headers: list[str]) -> dict[str, int]:
    """헤더 이름을 부분 일치로 찾아 {필드: 열 번호} 를 낸다."""
    flat = [re.sub(r"\s+", "", h).lower() for h in headers]
    found: dict[str, int] = {}
    used: set[int] = set()
    for key, hints in HINTS.items():
        for hint in hints:
            needle = re.sub(r"\s+", "", hint).lower()
            for index, header in enumerate(flat):
                if index in used or not header:
                    continue
                if needle in header:
                    found[key] = index
                    used.add(index)
                    break
            if key in found:
                break
    return found


def _cell(row: list[str], index: int | None) -> str:
    if index is None or index < 0 or index >= len(row):
        return ""
    return str(row[index]).strip()


def exposure_state(value: str) -> str:
    text = re.sub(r"\s+", "", value).lower()
    if not text:
        return ""
    for word in HIDDEN_WORDS:
        if text == word or (len(word) > 1 and word in text):
            return "미노출"
    for word in SHOWN_WORDS:
        if text == word or (len(word) > 1 and word in text):
            return "노출"
    return value.strip()


def sale_state(value: str) -> str:
    text = re.sub(r"\s+", "", value)
    if not text:
        return ""
    for word in SOLDOUT_WORDS:
        if word in text:
            return "품절"
    return "판매중"


def build_products(
    headers: list[str], body: list[list[str]], mapping: dict[str, int], cfg: dict
) -> tuple[list[fetcher.Product], dict]:
    """행을 상품으로 바꾼다. (상품 목록, 집계) 를 낸다."""
    out: list[fetcher.Product] = []
    stats = {"행": 0, "노출": 0, "미노출": 0, "품절": 0, "코드 없음": 0}
    for row in body:
        code = re.sub(r"[^0-9]", "", _cell(row, mapping.get("code")))
        if not code:
            stats["코드 없음"] += 1
            continue
        stats["행"] += 1
        prod = fetcher.Product(
            goods_no=code,
            url=fetcher.product_url(code, cfg.get("domain", "")),
            name=_cell(row, mapping.get("name")),
            price=re.sub(r"[^0-9]", "", _cell(row, mapping.get("price"))),
            brand=cfg.get("brand", "라포르"),
            seller=cfg.get("seller", "진바이오테크"),
            manufacturer=_cell(row, mapping.get("supplier")) or cfg.get("seller", ""),
            status="엑셀에서 읽음",
        )
        prod.exposed = exposure_state(_cell(row, mapping.get("expose"))) or "미확인"
        prod.sale_state = sale_state(_cell(row, mapping.get("sale")))
        if prod.sale_state == "품절":
            prod.availability = fields.OUT_OF_STOCK
            stats["품절"] += 1
        if prod.exposed == "노출":
            stats["노출"] += 1
        elif prod.exposed == "미노출":
            stats["미노출"] += 1
        prod.selected = prod.exposed != "미노출"
        out.append(prod)
    return out, stats


# ---------------------------------------------------------------------- 열 고르기 창


def ask_mapping(app, headers: list[str], guess: dict[str, int]) -> dict[str, int] | None:
    """열을 자동으로 찾지 못했을 때 사용자가 고르게 한다."""
    tk, ttk = app.tk, app.ttk
    win = tk.Toplevel(app.root)
    win.title("엑셀 열 고르기")
    win.geometry("560x420")
    win.transient(app.root)
    tk.Label(
        win,
        text="어떤 열이 무엇인지 골라 주세요. 상품코드만 필수이고 나머지는 비워 둘 수 있습니다.",
        justify="left",
        wraplength=520,
    ).pack(anchor="w", padx=12, pady=10)
    options = ["(사용 안 함)"] + [f"{i + 1}. {h}" for i, h in enumerate(headers)]
    vars_: dict[str, object] = {}
    body = ttk.Frame(win)
    body.pack(fill="x", padx=12)
    for row, (key, label) in enumerate(LABELS.items()):
        ttk.Label(body, text=("*" if key in REQUIRED else " ") + label).grid(
            row=row, column=0, sticky="e", padx=(0, 8), pady=4
        )
        var = tk.StringVar(value=options[guess[key] + 1] if key in guess else options[0])
        ttk.Combobox(body, textvariable=var, values=options, width=48, state="readonly").grid(
            row=row, column=1, sticky="w", pady=4
        )
        vars_[key] = var
    result: dict[str, int] = {}

    def confirm() -> None:
        picked: dict[str, int] = {}
        for key, var in vars_.items():
            value = var.get()
            if value and value != options[0]:
                picked[key] = int(value.split(".", 1)[0]) - 1
        if "code" not in picked:
            app.warn("엑셀 열 고르기", "상품코드 열은 반드시 골라야 합니다.")
            return
        result.update(picked)
        win.destroy()

    buttons = ttk.Frame(win)
    buttons.pack(fill="x", padx=12, pady=12)
    ttk.Button(buttons, text="확인", command=confirm).pack(side="left")
    ttk.Button(buttons, text="취소", command=win.destroy).pack(side="left", padx=6)
    win.grab_set()
    app.root.wait_window(win)
    return result or None


def import_dialog(app) -> tuple[list[fetcher.Product], dict, str] | None:
    """파일을 고르고 상품 목록을 만든다. (상품, 집계, 파일 경로)."""
    from tkinter import filedialog

    path = filedialog.askopenfilename(
        title="고도몰 관리자 상품 목록 엑셀 파일 고르기",
        filetypes=[("엑셀·CSV 파일", "*.xlsx *.xlsm *.csv *.tsv"), ("모든 파일", "*.*")],
    )
    if not path:
        return None
    try:
        headers, body = read_table(path)
    except Exception as exc:
        app.warn("엑셀 가져오기", f"파일을 읽지 못했습니다.\n{exc}")
        return None
    mapping = match_columns(headers)
    missing = [LABELS[k] for k in REQUIRED if k not in mapping]
    if missing or "name" not in mapping or "price" not in mapping:
        picked = ask_mapping(app, headers, mapping)
        if not picked:
            return None
        mapping = picked
    products, stats = build_products(headers, body, mapping, app.cfg)
    if not products:
        app.warn("엑셀 가져오기", "상품코드가 있는 행을 찾지 못했습니다.")
        return None
    return products, stats, path
